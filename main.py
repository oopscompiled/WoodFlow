import openpyxl
from datetime import datetime, timedelta
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Параметры
year = 2026
month = 3
template_file = 'шаблон.xlsx'  # Название файла шаблона
output_file = 'WoodFlow_журнал.xlsx'

# Загрузка шаблона
wb = openpyxl.load_workbook(template_file)
template_ws = wb.active

# Удаление существующих листов (кроме шаблона)
for sheet in wb.sheetnames[1:]:
    wb.remove(wb[sheet])

# Получение количества дней в месяце
if month == 12:
    last_day = (datetime(year + 1, 1, 1) - timedelta(days=1)).day
else:
    last_day = (datetime(year, month + 1, 1) - timedelta(days=1)).day

# Генерация листов для каждой смены (2/2)
current_date = datetime(year, month, 1)
shift_counter = 0  # 0-1: рабочие дни, 2-3: выходные
sheet_count = 0

while current_date.month == month:
    if shift_counter < 2:  # Рабочие дни
        date_str = current_date.strftime("%d.%m.%Y")
        
        # Копирование шаблона
        new_ws = wb.copy_worksheet(template_ws)
        new_ws.title = date_str
        
        # Редактирование шаблона под дату
        new_ws['A1'] = f"Посменный журнал: {date_str}"
        
        sheet_count += 1
        current_date += timedelta(days=1)
    else:  # Выходные дни
        current_date += timedelta(days=1)
    
    shift_counter = (shift_counter + 1) % 4  # Цикл 0,1,2,3

# Сохранение файла
wb.save(output_file)
print(f"Файл создан: {output_file}")
print(f"Листов с датами: {sheet_count}")